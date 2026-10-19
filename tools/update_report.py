from pathlib import Path

from docx import Document


SOURCE = Path("Database&Docs/CNLTHD.docx")
OUTPUT = Path("Database&Docs/CNLTHD_DaChinhSua.docx")


REPLACEMENTS = {
    "Trong trạng thái repository hiện tại chưa có test project": (
        "Repository hiện có project TaskFlow.Tests. Bộ kiểm thử tự động đã chạy thành công 7 trường hợp, "
        "gồm kiểm tra mật khẩu sai, từ chối DEMO_PASSWORD_HASH, đăng ký luôn nhận vai trò Member và các "
        "ràng buộc Data Annotations tiêu biểu. Kiểm thử tích hợp trực tiếp với SQL Server và migration trên "
        "database mới vẫn cần được thực hiện riêng trên môi trường chạy thật."
    ),
    "Kết quả kiểm thử cho thấy các project TaskFlow.Domain": (
        "Lệnh dotnet test biên dịch toàn bộ solution và chạy thành công 7/7 kiểm thử. Tuy nhiên, NuGet vẫn "
        "cảnh báo phụ thuộc bắc cầu System.Security.Cryptography.Pkcs 9.0.11 có lỗ hổng đã công bố; vì vậy "
        "báo cáo không coi kết quả build là bằng chứng rằng dependency đã hoàn toàn an toàn."
    ),
    "Đối với chức năng đăng ký, AuthController": (
        "Đối với chức năng đăng ký, AuthController kiểm tra dữ liệu đầu vào và yêu cầu mật khẩu có độ dài tối "
        "thiểu 6 ký tự. AuthService kiểm tra tính duy nhất của username và email, sau đó luôn lấy vai trò Member "
        "từ cơ sở dữ liệu; RegisterRequest không nhận RoleId nên client không thể tự chọn Admin hoặc Manager. "
        "Mật khẩu được băm trước khi lưu."
    ),
    "Về mặt kỹ thuật, backend đã được build với target framework net10.0. Tuy nhiên, repository hiện chưa có project kiểm thử": (
        "Backend sử dụng target framework net10.0 và repository đã có project TaskFlow.Tests. Lần chạy gần nhất "
        "đạt 7/7 kiểm thử, bao phủ một phần xác thực và validation. Các ca HTTP end-to-end cho 401/403 theo từng "
        "vai trò và kiểm thử transaction trên SQL Server vẫn là phần cần bổ sung minh chứng runtime."
    ),
    "Về khả năng vận hành, backend đã được build thành công trong lần kiểm tra trước": (
        "Về khả năng vận hành, solution đã build và chạy 7/7 kiểm thử tự động thành công. Kết quả này chứng minh "
        "các ca đã viết, nhưng chưa thay thế kiểm thử end-to-end bằng Postman/Swagger và SQL Server cho toàn bộ "
        "ma trận quyền Admin, Manager, Member."
    ),
    "Thứ hai, quyền đăng ký tài khoản cần được kiểm soát chặt chẽ": (
        "Quyền đăng ký công khai đã được kiểm soát: RegisterRequest không có RoleId và AuthService luôn gán vai "
        "trò Member lấy từ bảng Roles. Việc cấp Admin hoặc Manager không nằm trong API đăng ký công khai và phải "
        "được thực hiện bằng quy trình quản trị riêng."
    ),
    "Thứ ba, cơ chế xác minh mật khẩu dành cho dữ liệu demo": (
        "Nhánh đăng nhập bằng mật khẩu mẫu đã được loại bỏ khỏi PasswordVerifier. Chuỗi DEMO_PASSWORD_HASH trong "
        "dữ liệu mẫu không được chấp nhận; tài khoản mẫu chỉ đăng nhập được sau khi được thay bằng chuỗi PBKDF2 hợp lệ."
    ),
    "Thứ tư, hệ thống nên bổ sung Refresh Token Rotation": (
        "Access token hiện hết hạn sau 8 giờ theo cấu hình Jwt:ExpiresHours và được kiểm tra lifetime với ClockSkew "
        "một phút. Phạm vi hiện tại chưa triển khai refresh token, danh sách thu hồi hoặc logout phía server; do đó "
        "token đã cấp chỉ mất hiệu lực khi hết hạn hoặc khi thay khóa ký. Đây là giới hạn cần nêu rõ khi triển khai production."
    ),
    "Tuy nhiên, quy trình tạo project hiện lưu thông tin project": (
        "Quy trình tạo project lưu project và membership của owner trong một transaction EF Core. Hai lần "
        "SaveChangesAsync được commit cùng nhau; nếu bước thêm owner thất bại, transaction không được commit và "
        "toàn bộ thay đổi được hoàn tác."
    ),
    "Nhìn chung, TaskFlow đã có một số biện pháp bảo đảm tính nhất quán dữ liệu": (
        "TaskFlow đã bổ sung transaction cho thao tác tạo project nhiều bước và duy trì các ràng buộc khóa chính, "
        "khóa ngoại ở database. Retry policy cho lỗi SQL Server tạm thời, refresh token và giám sát nâng cao vẫn "
        "chưa nằm trong phạm vi hiện tại."
    ),
    "Để đánh giá độ chính xác của cơ chế Authentication và Authorization, nhóm nên xây dựng": (
        "Nhóm đã xây dựng kiểm thử tự động cho mật khẩu sai, chuỗi DEMO_PASSWORD_HASH, đăng ký mặc định Member "
        "và validation DTO. Các trường hợp còn cần bổ sung gồm token sai/hết hạn, ma trận quyền Admin/Manager/Member, "
        "người ngoài project, assignment sai phạm vi và kiểm thử HTTP end-to-end."
    ),
    "Việc triển khai các trường hợp này thành integration test": (
        "Các kiểm thử hiện tại chủ yếu ở mức unit/service với EF Core InMemory. Integration test qua HTTP và SQL "
        "Server vẫn cần bổ sung để kiểm tra đồng thời middleware, authorization, transaction và truy vấn thực tế."
    ),
    "Các nội dung trên hiện là những hạng mục cần được bổ sung": (
        "Các ca xác thực và validation cốt lõi đã có kiểm thử tự động; phần còn thiếu được giới hạn rõ ở kiểm thử "
        "HTTP, SQL Server, ma trận phân quyền đầy đủ và số liệu observability runtime."
    ),
    "RegisterRequest: Tiếp nhận dữ liệu đăng ký người dùng mới": (
        "RegisterRequest tiếp nhận Username, Email, Password, FullName và AvatarUrl. DTO dùng Required, "
        "StringLength, EmailAddress và Url; không có RoleId. AuthService luôn gán vai trò Member cho đăng ký công khai."
    ),
    "AssignTaskRequest: Tiếp nhận danh sách mã định danh UserIds": (
        "AssignTaskRequest tiếp nhận một UserId nguyên dương. Service chỉ cho phép phân công người dùng đang hoạt "
        "động và là thành viên của project chứa task."
    ),
    "Nhóm đã tiến hành kiểm tra hoạt động của cơ chế Model Binding và Validation": (
        "Cơ chế Model Binding và Validation được đối chiếu bằng mã nguồn, build và kiểm thử tự động; các ảnh thực "
        "nghiệm Postman/Swagger cần được chụp lại theo checklist minh chứng đi kèm báo cáo."
    ),
    "Khi gửi request đến POST /api/auth/register với email sai định dạng": (
        "Khi gửi POST /api/auth/register với email sai định dạng hoặc mật khẩu ngắn hơn 6 ký tự, Data Annotations "
        "kết hợp [ApiController] trả 400 trước khi action thực thi. Chuỗi chỉ gồm khoảng trắng tiếp tục được Controller kiểm tra thủ công."
    ),
    "Các kết quả trên cho thấy hệ thống TaskFlow đã thiết lập": (
        "Kết quả đối chiếu cho thấy TaskFlow có model binding theo kiểu mạnh, Data Annotations trên Request DTO, "
        "[ApiController] và kiểm tra liên trường tại Controller. Hai kiểm thử validation tiêu biểu đã chạy thành công; "
        "các ảnh HTTP thực tế vẫn cần bổ sung theo checklist."
    ),
    "Ngăn ngừa lỗ hổng Over-posting / Mass Assignment": (
        "Ngăn ngừa Over-posting/Mass Assignment: các API nhận Request DTO thay vì entity. RegisterRequest không có "
        "RoleId; CreateTaskRequest không có CreatedById; CreateCommentRequest và CreateTaskHistoryRequest không có "
        "UserId. Các giá trị định danh người thực hiện được lấy từ claim JWT."
    ),
    "Tách rời quy tắc kiểm tra bằng thư viện chuyên dụng": (
        "Các quy tắc cấu trúc hiện dùng Data Annotations trực tiếp trên DTO, còn quy tắc liên trường và nghiệp vụ "
        "được kiểm tra tại Controller/service. FluentValidation có thể được cân nhắc khi số lượng quy tắc phức tạp tăng."
    ),
    "Kiểm tra tính logic của các trường thời gian": (
        "Controller hiện kiểm tra ngày kết thúc không trước ngày bắt đầu đối với project và task; đồng thời kiểm tra "
        "số giờ không âm. Quy tắc hạn chót không nằm trong quá khứ chưa được áp dụng vì có thể cản trở việc nhập dữ liệu lịch sử."
    ),
    "Để nâng cao độ tin cậy và đảm bảo tính đúng đắn của tầng Model Binding": (
        "Project TaskFlow.Tests đã có kiểm thử Validator.TryValidateObject cho email, mật khẩu, ID dương và số giờ "
        "không âm. Các ca biên còn lại và integration test với WebApplicationFactory vẫn cần được mở rộng."
    ),
    "Model Binding và Validation là lớp kiểm tra dữ liệu đầu vào": (
        "Model Binding và Validation là lớp kiểm tra dữ liệu đầu vào của TaskFlow. Request DTO dùng Data Annotations, "
        "[ApiController] tự động trả ValidationProblemDetails cho lỗi binding/annotation, còn Controller và service xử "
        "lý điều kiện liên trường, tham chiếu và quyền truy cập."
    ),
    "Định dạng ValidationProblemDetails (RFC 7807) giúp frontend": (
        "ValidationProblemDetails được dùng cho lỗi model binding và Data Annotations. Một số nhánh BadRequest thủ "
        "công và middleware ngoại lệ vẫn trả cấu trúc khác, vì vậy chuẩn hóa toàn bộ lỗi 400/401/403/404/500 vẫn là hạng mục cần hoàn thiện."
    ),
    "Dữ liệu minh họa còn gồm ba tài khoản": (
        "Dữ liệu minh họa gồm ba tài khoản, một dự án, hai thành viên và một công việc được phân công. "
        "DEMO_PASSWORD_HASH chỉ là giá trị giữ chỗ và PasswordVerifier chủ động từ chối giá trị này; cần thay bằng "
        "chuỗi PBKDF2 hợp lệ trước khi kiểm thử đăng nhập."
    ),
    "CreateTaskRequest gồm ProjectId, CreatedById": (
        "CreateTaskRequest gồm ProjectId, StatusId, PriorityId, tiêu đề, mô tả, ngày bắt đầu, thời hạn và số giờ. "
        "CreatedById không nhận từ body; backend lấy người tạo từ claim NameIdentifier của JWT."
    ),
    "Khi tạo công việc, Controller truyền thêm định danh lấy từ token": (
        "Khi tạo công việc, Controller lấy định danh người dùng từ JWT và service dùng trực tiếp giá trị này làm "
        "CreatedById, nhờ đó client không thể giả mạo người tạo trong payload."
    ),
    "Khi tạo bình luận hoặc lịch sử, TaskId trong body": (
        "TaskId được lấy từ route và UserId được lấy từ claim JWT; body bình luận chỉ chứa Content, còn body lịch sử "
        "chứa Action, OldValue và NewValue. Nội dung bắt buộc được kiểm tra bằng Data Annotations và Controller."
    ),
    "Các Controller bình luận và lịch sử trong mã nguồn": (
        "CommentsController và TaskHistoriesController yêu cầu xác thực. Service chỉ trả dữ liệu khi người dùng là "
        "Admin, chủ project hoặc assignee còn thuộc project; lịch sử chỉ được tạo bởi Admin/Manager có quyền quản lý, "
        "còn bình luận chỉ tác giả, Admin hoặc chủ project mới được sửa/xóa."
    ),
    "Tương tự nhóm bình luận và lịch sử, hai Controller danh mục": (
        "StatusesController và PrioritiesController yêu cầu đăng nhập cho thao tác đọc; POST, PUT và DELETE chỉ cho "
        "vai trò Admin. Fallback policy toàn cục cũng yêu cầu người dùng xác thực, ngoại trừ endpoint được đánh dấu AllowAnonymous."
    ),
    "Migration được sử dụng để quản lý thay đổi cấu trúc cơ sở dữ liệu theo mô hình EF Core": (
        "Migration được quản lý trong TaskFlow.Infrastructure/Migrations. Repository hiện có migration "
        "20261011132137_InitialCreate, file Designer và TaskFlowDbContextModelSnapshot; lệnh áp dụng trên database "
        "trống cần được chụp làm minh chứng runtime."
    ),
    "Các Controller dự án, công việc và phân công sử dụng [Authorize]": (
        "Fallback policy yêu cầu xác thực cho toàn bộ endpoint, trừ AuthController và HelloController được đánh dấu "
        "[AllowAnonymous]. Project, Task, Assignment, Comment, History, Status và Priority đều có quy tắc vai trò hoặc phạm vi tài nguyên tương ứng."
    ),
    "Trong phạm vi mã nguồn đã đối chiếu, các Controller bình luận": (
        "Sau khi rà soát, API bình luận, lịch sử, trạng thái và mức độ ưu tiên đã được bảo vệ. Quyền đọc Comment/History "
        "được giới hạn theo task; quyền ghi History theo quản lý project; quyền thay đổi Status/Priority chỉ dành cho Admin."
    ),
}


def replace_paragraph(paragraph, new_text):
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


document = Document(SOURCE)
matched = set()
for paragraph in all_paragraphs(document):
    text = paragraph.text.strip()
    for prefix, replacement in REPLACEMENTS.items():
        if text.startswith(prefix):
            replace_paragraph(paragraph, replacement)
            matched.add(prefix)
            break

table = document.tables[19]
table.rows[10].cells[1].text = "1 migration + 1 model snapshot"
table.rows[10].cells[2].text = "Schema EF Core đã được phiên bản hóa trong Infrastructure/Migrations"

auth_results = document.tables[21]
auth_results.rows[2].cells[3].text = "Tự động: PasswordVerifierTests và AuthServiceTests"
auth_results.rows[3].cells[3].text = "Cần chụp lại HTTP 401"
auth_results.rows[4].cells[3].text = "Cần chụp lại HTTP 401"
auth_results.rows[5].cells[3].text = "Cần chụp lại HTTP 403"
auth_results.rows[6].cells[3].text = "Cần chụp lại HTTP 201"

validation_results = document.tables[23]
validation_results.rows[1].cells[1].text = "15 Request DTO"
validation_results.rows[2].cells[1].text = "[FromBody], [FromRoute] và binding mặc định"
validation_results.rows[4].cells[1].text = "RFC 7807 cho lỗi model; lỗi thủ công chưa hoàn toàn đồng nhất"
validation_results.rows[5].cells[1].text = "Build thành công; 7/7 kiểm thử đạt"

comments_api = document.tables[35]
comments_api.rows[1].cells[2].text = "Tạo bình luận; taskId từ route, UserId từ JWT, body chứa Content."
comments_api.rows[6].cells[2].text = "Tạo lịch sử; taskId từ route, UserId từ JWT, body chứa Action, OldValue, NewValue."

statuses_api = document.tables[36]
for row_index in (1, 4, 5, 6, 9, 10):
    statuses_api.rows[row_index].cells[2].text += " Chỉ Admin."
for row_index in (2, 3, 7, 8):
    statuses_api.rows[row_index].cells[2].text += " Yêu cầu đăng nhập."

appendix_index = next(i for i, p in enumerate(document.paragraphs) if p.text.strip().startswith("Phụ lục A"))
anchor = document.paragraphs[appendix_index]
anchor.insert_paragraph_before("2. Kết quả đối chiếu mã nguồn sau góp ý", style="Heading 2")
anchor.insert_paragraph_before(
    "Mã nguồn đã được rà soát lại theo các nhận xét. Các thay đổi chính gồm loại bỏ nhánh mật khẩu demo; "
    "không nhận RoleId, CreatedById hoặc UserId nhạy cảm từ body; bổ sung fallback authorization; kiểm tra "
    "phạm vi truy cập Comment/History; giới hạn thay đổi Status/Priority cho Admin; dùng transaction khi tạo project; "
    "và ghi log trước/sau middleware kèm TraceId, status code, thời gian xử lý."
)
anchor.insert_paragraph_before(
    "Các Request DTO đã có Data Annotations và solution có TaskFlow.Tests. Lần chạy gần nhất đạt 7/7 kiểm thử. "
    "Migration hiện có một phiên bản InitialCreate cùng model snapshot. Những kết quả này được ghi theo đúng "
    "mã nguồn; ảnh Postman, Swagger, log và SQL Server chưa được tự suy diễn là đã có."
)
anchor.insert_paragraph_before(
    "Giới hạn còn lại gồm chưa có refresh token/token revocation phía server, chưa chuẩn hóa mọi phản hồi lỗi về "
    "một schema duy nhất, chưa có integration test SQL Server và còn cảnh báo bảo mật ở phụ thuộc bắc cầu "
    "System.Security.Cryptography.Pkcs 9.0.11. Danh sách ảnh cần chụp được tách thành file GhiChu_AnhMinhChung.md."
)

missing = set(REPLACEMENTS) - matched
if missing:
    raise RuntimeError("Unmatched report paragraphs: " + "; ".join(sorted(missing)))

document.save(OUTPUT)
print(OUTPUT)
